"""Read-only reconciliation of the supplied FC2B export, not a new backtest.

No rule selection, indicator reconstruction, performance search or account run.
One explicitly labeled fixed-ledger NQ cost rescore; original files preserved.
"""
from pathlib import Path
import hashlib
import json
import re
import sys
from importlib.metadata import version

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from resampling import max_dd

ROOT=Path(__file__).resolve().parents[1]
INPUT=ROOT/'agents/codex/inputs/fc2b'
OUT=ROOT/'analysis/outputs/codex_fc2b_report_audit'
OVERVIEW={2017:(83,616.),2018:(83,2060.5),2019:(97,958.),2020:(83,1281.5),
          2021:(85,3436.),2022:(95,5244.5),2023:(82,468.),2024:(91,2752.5),
          2025:(89,3841.),2026:(68,1468.)}


def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def dump(p,x):Path(p).write_text(json.dumps(x,indent=2,default=str)+'\n')
def money(x):return ('-$' if x<0 else '$')+f'{abs(x):,.2f}'


def main():
    OUT.mkdir(parents=True,exist_ok=True)
    wb=load_workbook(INPUT/'report.xlsx',read_only=True,data_only=True)
    settings={r[0]:r[1] for r in wb['Settings'].iter_rows(values_only=True) if r[0] is not None}
    pd.DataFrame(list(settings.items()),columns=['setting','value']).to_csv(OUT/'settings.csv',index=False)
    summaries={r[0]:list(r[1:4]) for r in wb['Strategy Analysis'].iter_rows(values_only=True) if r[0] is not None}
    d=pd.read_excel(INPUT/'report.xlsx',sheet_name='List of Trades',header=2)
    e=d[d.Type.astype(str).str.startswith('Entry')].reset_index(drop=True)
    x=d[d.Type.astype(str).str.startswith('Exit')].reset_index(drop=True)
    assert len(e)==len(x) and len(e)>0
    assert e['Trade #'].is_unique
    t=pd.DataFrame(dict(number=e['Trade #'].astype(int),entry_time=pd.to_datetime(e.Date),exit_time=pd.to_datetime(x.Date),
        side=np.where(e.Type.str.contains('Long'),1,-1),entry=e.Price,exit=x.Price,contracts=e.Contracts,
        gross=e['Profit ($)'],mae=e['Drawdown ($)'],mfe=e['Run-up ($)'],exit_signal=x.Signal))
    assert t.contracts.eq(1).all() and settings['Point Value']=='$20'
    assert (t.exit_time>=t.entry_time).all()
    assert ((t.side==1)==x.Type.str.contains('Long')).all()
    assert np.allclose(t.gross,(t['exit']-t.entry)*t.side*20*t.contracts,rtol=0,atol=1e-7)
    assert float(summaries['Total # of Trades'][0])==len(t)
    assert np.isclose(float(summaries['Net Profit'][0]),t.gross.sum())
    assert np.isclose(float(summaries['Gross Profit'][0]),t.gross[t.gross>0].sum())
    assert np.isclose(float(summaries['Gross Loss'][0]),t.gross[t.gross<0].sum())
    assert float(summaries['Slippage Paid'][0])==float(summaries['Commission Paid'][0])==0.
    assert np.isclose(abs(float(summaries['Max Close To Close Drawdown'][0])),max_dd(t.gross))
    expected=dict(ORStartTime=830,ORMinutes=30,ExitTime=1458,ZMin=1.,ZLookback=60,MinWarmup=20,
                  SLpoints=100,TrailPts=150,BreakEvenR=.66,FixedContracts=1)
    source=(INPUT/'source.txt').read_text()
    inputblock=re.search(r'inputs:(.*?)variables:',source,re.I|re.S).group(1)
    for name,value in expected.items():
        code=float(re.search(r'\b'+re.escape(name)+r'\s*\(\s*([0-9.]+)',inputblock,re.I).group(1))
        assert code==value==float(settings[name]),name
    # This does not change any price/order sequence or execute a strategy.
    t['repo_cost_rescore_NQ']=t.gross-17.25
    t.to_csv(OUT/'trades.csv',index=False)
    metrics=[]
    periods=[('full',t),('2017-22',t[t.entry_time.dt.year<=2022]),
             ('2018-22',t[t.entry_time.dt.year.between(2018,2022)]),
             ('2023+',t[t.entry_time.dt.year>=2023])]+[(str(y),t[t.entry_time.dt.year==y]) for y in OVERVIEW]
    for period,part in periods:
        for side,v in [('all',0),('long',1),('short',-1)]:
            z=part if v==0 else part[part.side==v]
            for label,column in [('export_gross','gross'),('analytic_NQ_17.25','repo_cost_rescore_NQ')]:
                net=z[column]
                metrics.append(dict(period=period,side=side,accounting=label,trades=len(z),total=float(net.sum()),
                    average=float(net.mean()),win=float(net.gt(0).mean()),
                    profit_factor=float(net[net>0].sum()/-net[net<0].sum()),closed_trade_max_dd=max_dd(net)))
    m=pd.DataFrame(metrics);m.to_csv(OUT/'fixed_ledger_metrics.csv',index=False)
    years=[]
    for y,(count,pdfnet) in OVERVIEW.items():
        z=t[t.entry_time.dt.year==y]
        years.append(dict(year=y,export_trades=len(z),overview_trades=count,export_gross_NQ=float(z.gross.sum()),
                          overview_claimed_net_MNQ=pdfnet,overview_net_NQ_equivalent=pdfnet*10,
                          export_gross_less_overview_NQ_equivalent=float(z.gross.sum()-pdfnet*10)))
    years=pd.DataFrame(years);years.to_csv(OUT/'overview_year_comparison.csv',index=False)
    # Price observations only, not a reconstructed stop-fill counterfactual.
    spots=[];price_sources={}
    for r in t[t.gross.lt(-2000)].itertuples():
        path=ROOT/f'data/NQ_1min_{r.entry_time.year}.csv';bars=pd.read_csv(path)
        day=int(r.entry_time.strftime('%Y%m%d'));j=r.entry_time.hour*100+r.entry_time.minute
        b=bars[(bars.date==day)&(bars.time==j)].iloc[0]
        nextj=r.exit_time.hour*100+r.exit_time.minute
        bnext=bars[(bars.date==day)&(bars.time==nextj)].iloc[0]
        level=r.entry-r.side*100
        crossed=bool(b.low<=level if r.side==1 else b.high>=level)
        spots.append(dict(number=r.number,entry_time=str(r.entry_time),exit_time=str(r.exit_time),side=r.side,
             entry=r.entry,nominal_stop=level,entry_minute_open=float(b.open),entry_minute_high=float(b.high),
             entry_minute_low=float(b.low),entry_minute_close=float(b.close),
             nominal_stop_crossed_in_entry_minute=crossed,reported_exit=r.exit,
             exit_minute_open=float(bnext.open),gross_NQ=r.gross,loss_points=-r.gross/20))
        assert r.entry==b.open and r.exit==bnext.open and crossed
        price_sources[str(path.relative_to(ROOT))]=sha(path)
    pd.DataFrame(spots).to_csv(OUT/'entry_minute_loss_spotchecks.csv',index=False)
    prior=np.r_[0,t.gross.cumsum().to_numpy()[:-1]]
    closed_peaks=np.maximum.accumulate(prior)
    mae_from_closed_peak=float(np.max(closed_peaks-(prior+t.mae.to_numpy())))
    gross=m[(m.period=='full')&(m.side=='all')&(m.accounting=='export_gross')].iloc[0]
    cost=m[(m.period=='full')&(m.side=='all')&(m.accounting=='analytic_NQ_17.25')].iloc[0]
    onepoint=int(np.isclose(t.gross,20).sum())
    summary=dict(trades=len(t),long_trades=int(t.side.eq(1).sum()),short_trades=int(t.side.eq(-1).sum()),
       first_entry=str(t.entry_time.min()),last_exit=str(t.exit_time.max()),gross_NQ=float(t.gross.sum()),
       win_rate=float(t.gross.gt(0).mean()),export_profit_factor=float(gross.profit_factor),
       export_strategy_drawdown_NQ=abs(float(summaries['Max Strategy Drawdown'][0])),
       export_closed_drawdown_NQ=abs(float(summaries['Max Close To Close Drawdown'][0])),
       closed_peak_plus_trade_MAE_drawdown_NQ=mae_from_closed_peak,max_trade_MAE_NQ=float(-t.mae.min()),
       largest_loss_NQ=float(t.gross.min()),one_point_gross_winners=onepoint,
       entries_all_0901=bool(t.entry_time.dt.strftime('%H:%M').eq('09:01').all()),
       intraday_all=bool((t.entry_time.dt.date==t.exit_time.dt.date).all()),
       end_of_day_fallbacks=int(t.exit_signal.eq('End of Day Exit').sum()),
       no_commission_or_slippage=True,source_parameters_match_export=True,
       not_new_backtest=True,new_strategy_trials=0,analytic_fixed_ledger_cost_rescores=1)
    dump(OUT/'summary.json',summary)
    rows=['| '+str(r.year)+' | '+str(r.export_trades)+' | '+str(r.overview_trades)+' | '+money(r.export_gross_NQ)+' | '+money(r.overview_net_NQ_equivalent)+' |' for r in years.itertuples()]
    report=f'''# FC 2B export audit: WATCH — inputs match, but this is a different run from the overview

Summary: The supplied @NQ export has {len(t)} trades and {money(gross.total)} gross, with no commission or slippage.
Its settings match the supplied code inputs; it is not the overview's 856-trade, cost-adjusted MNQ report.
The largest loss crosses the nominal stop in the entry minute and exits at the following minute's open.
Analytic repo-cost rescore stays profitable, but does not verify MNQ fees, signal parity or prospective performance.
Obtain the overview's exact source ledger, Volume/Ticks observations and forward log before promotion or new variants.

Plan: not applicable — existing-export reconciliation, not a new strategy/backtest or preregistered return trial.
Script: `analysis/codex_fc2b_report_audit.py`; inputs: `agents/codex/inputs/fc2b/`.
Outputs: `analysis/outputs/codex_fc2b_report_audit/`. New strategy trials **0**; analytic cost rescores **1**.

## Rule and settings as supplied

30-minute opening drive on 1-minute @NQ: ORStart 830, ORMinutes 30, ExitTime 1458; ZMin 1, ZLookback 60,
MinWarmup 20; stop 100 points, trail 150, break-even 0.66R, fixed one contract. All ten input values match the code.
Settings: point value $20, one-minute compression, Classic, Bar Magnifier Disabled, initial capital $100,000,
commission "{settings['Commission']}", slippage "{settings['Slippage']}".
The export does not establish the chart's Build Volume On setting, IOG checkbox, exact session template or real
broker/order settings. "Automate Trading (H): Off" is recorded; no live state or account activity is inferred.

Settings Start Date {settings['Start Date']}, End Date {settings['End Date']}; actual trades
{summary['first_entry']} through {summary['last_exit']}. Source timestamps are naive/chart-local; no export-time
timezone conversion is invented. Input source contains live-value comments, which authorize no deployment.

## Reconciled export results

All {len(t)} entry/exit pairs are complete, one contract each. Every trade's reported profit exactly equals
side×(exit−entry)×$20. Summary counts, gross profits/losses, total and closed-equity DD reconcile to the ledger.
Longs {summary['long_trades']}, shorts {summary['short_trades']}; all entries stamped 09:01, all exits same date.

| Accounting | Trades | Total/NQ | Average/NQ | Win | PF | Closed-trade DD/NQ |
|---|---:|---:|---:|---:|---:|---:|
| Export, no costs | {len(t)} | {money(gross.total)} | {money(gross.average)} | {gross.win:.2%} | {gross.profit_factor:.4f} | {money(gross.closed_trade_max_dd)} |
| Analytic $17.25/NQ debit | {len(t)} | {money(cost.total)} | {money(cost.average)} | {cost.win:.2%} | {cost.profit_factor:.4f} | {money(cost.closed_trade_max_dd)} |

The cost debit is the repo assumption $7.25 commission plus one $5 tick each market/stop fill. It changes no fills,
signals, orders or stop anchors. It is an arithmetic accounting sensitivity, not an execution simulation, fresh
edge test or verified MNQ result. Full period/year/side ledger summaries are saved. No new bootstrap, strategy
backtest, placebo, full minute replay or funded bench was run for this audit.

## Why it does not reproduce the overview

The overview claims 856 trades from 2017-02-02, 416 longs/440 shorts, and $22,126 net per MNQ.
This export has 826 trades from 2017-05-30, 404 longs/422 shorts. All 30 missing trades occur in 2017; 2018–2026
annual counts match. Shorter initial history/warm-up could explain the early count difference, but is unproven.
Later annual profits also differ; some PDF net profits exceed this export's gross profit at the same point value.
Fees and missing 2017 trades alone cannot reconcile that. Different entries/prices/data/settings history or
report postprocessing remain possible causes; no cause is selected without paired original ledgers.

| Year | Export trades | Overview trades | Export gross/NQ | Overview claimed net ×10 |
|---|---:|---:|---:|---:|
{chr(10).join(rows)}

In 2020 export gross $12,330 is below overview NQ-equivalent net $12,815; in 2025 gross $37,865 is below net $38,410.
The overview may legitimately apply costs outside MultiCharts, so its claim is not disproved merely by the raw
zero-cost Settings. The exact source report and accounting transformation are still required.
Raw max win $8,870 / max loss −$3,835 become $887 / −$383.50 at MNQ point value; the overview's $885.50 / −$385
extrema are each consistent with a $1.50 micro debit, but that observation does not reconcile the annual ledgers.

## Execution and risk findings

- **Entry-minute exposure:** 2020-03-03 short enters 8780.25 in bar 09:01; nominal stop 8880.25. Owned entry-minute
  high 8978 crosses that level. The reported exit is 8972 at bar 09:02 open: **−191.75 points / −$3,835 per NQ**.
  The second overshoot, 2024-09-04, is a long losing 119.5 points / −$2,390. Both reported entries/exits match the
  owned minute opens; these are spot checks, not whole-strategy/data parity. See `entry_minute_loss_spotchecks.csv`.
  Under the source's documented IOG-off configuration, its first next-bar stop is only generated after the entry
  bar closes. No hypothetical protective fill or corrected-stop P&L was substituted into the ledger.
- **Moving stops:** BE/trail updates use completed bar extrema, act next bar and can already be on the wrong side
  of market after a reversal. No stop-rejection/fallback evidence is supplied. Replaying them retroactively within
  the triggering bar would be a different implementation.
- **Volume semantics:** the code weights `Volume`; official MultiCharts intraday documentation defines this as
  up-tick volume/count. Our `NQ_Bar_Exporter.txt` exports `Ticks`, the total volume/count, under CSV `volume`.
  Do not claim signal parity from total-volume OHLCV alone or silently replace the input. Obtain paired
  Time/Volume/Ticks/barDelta and daily ratio/z observations with actual Build Volume On settings.
- **Risk definitions:** export strategy DD **{money(summary['export_strategy_drawdown_NQ'])}**, closed DD
  **{money(summary['export_closed_drawdown_NQ'])}**, max trade MAE **{money(summary['max_trade_MAE_NQ'])}**, all per NQ
  before fees. A closed-equity-peak plus trade-MAE calculation gives **{money(mae_from_closed_peak)}**; it is not
  the same statistic as MC Max Strategy Drawdown, which can reflect different open-equity peak handling.
  No drawdown-definition mismatch is papered over, and neither number is an Apex breach probability.
- **Micro costs:** **{onepoint} trades ({onepoint/len(t):.1%}) earn exactly +$20 gross/NQ = +$2 at MNQ point value**.
  Actual micro round-trip costs above $2 would turn these outcomes negative; costs exactly $2 make them flat.
  A higher fee-adjusted losing-trade count need not eliminate positive overall expectancy. Recalculate using
  actual micro fees/fills; dividing full-contract net by 10 is not proof of micro net or its win rate.
- **Trading clock:** all positions are intraday in the supplied list; {summary['end_of_day_fallbacks']} use the
  platform End of Day Exit label rather than the explicit time order. Session/holiday handling needs inspection
  before relying on SetExitOnClose. Bar 09:01 is a close stamp: its open corresponds to 09:00 CT under owned data.

## Evidence status and next step

This export verifies a profitable historical gross ledger and matching numerical inputs. It does not supply the
overview's exact ledger, confirm Build Volume On/IOG settings, reproduce entry signals, validate live fills,
establish prospectively frozen performance or earn the repo's B grade. The earlier 211-trade holdout applies to
the original range-stop rule; current 100/150 exits were selected on the full sample, as the overview discloses.
Its unfinished 30-signal forward test and later-period selection history remain material.

The PDF's 2024–26 statement of 239 trades still differs from its own yearly totals 248. The new workbook also
has 248 trades in those years, so that paragraph likely describes an older window; exact run identity is needed.

Next smallest steps: obtain the exact 856-trade report/ledger behind the PDF and its cost calculation; obtain a
small paired Volume/Ticks/ratio/z export and the existing forward log; reconcile the supplied version first.
Any entry-bar protection, market fallback or total-volume substitution must be a separately frozen amendment,
with before/after parity and a declared trial status. Do not tune another filter or treat retrospective 2023+
results as a clean holdout. Portfolio addition/replacement and timed Apex risk come only after those gates.

## Reproduction and sources

```bash
/workspace/venvs/thebeginning/bin/python analysis/codex_fc2b_report_audit.py
```

User inputs are preserved byte-for-byte with hashes; two existing price CSVs were read only for the specific
overshoot observations. No data cache, strategy engine, rank, trading plan, live/VPS/broker file or API key was
changed/accessed. Source-document instructions are treated as evidence, not orders to trade or deploy.

Primary platform references:
[Volume](https://www.multicharts.com/trading-software/index.php?title=Volume),
[Ticks](https://www.multicharts.com/trading-software/index.php?title=Ticks).
Uploaded inputs: `report.xlsx`, `source.txt`, `overview.pdf`, plus its text extraction in the input directory.

Tell Claude: review `agents/codex/results/fc2b_report_audit.md`; 826 zero-cost NQ trades do not reconcile to the
856-trade net-MNQ overview, entry-minute stop exposure is demonstrated, and volume/input parity remains open.
'''
    (ROOT/'agents/codex/results/fc2b_report_audit.md').write_text(report)
    dump(OUT/'reproduction.json',dict(python=sys.executable,packages={p:version(p) for p in ['numpy','pandas','openpyxl']},
       input_sha256={str(p.relative_to(ROOT)):sha(p) for p in INPUT.iterdir() if p.is_file()},
       price_source_sha256=price_sources,code_sha256={str(Path(__file__).relative_to(ROOT)):sha(__file__),
           'analysis/resampling.py':sha(ROOT/'analysis/resampling.py')},
       operation='existing-export audit; one arithmetic cost rescore; no strategy backtest',new_strategy_trials=0))
    print(json.dumps(summary,indent=2))


if __name__=='__main__':main()
