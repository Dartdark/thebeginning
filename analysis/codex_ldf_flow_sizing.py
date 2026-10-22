"""Frozen exploratory LDF sizing screen. Run inputs first, then design only.

Plan agents/codex/plans/ldf_flow_sizing.md, amended before outcomes (6ca8950).
No 2023+ candidate features, strategy returns, funded simulation or cache creation.
"""
from __future__ import annotations

import argparse
import copy
import hashlib
import io
import json
import pickle
import subprocess
import sys
import tempfile
from datetime import datetime, timezone
from pathlib import Path

import exchange_calendars as xc
import numpy as np
import openpyxl
import pandas as pd

sys.path.insert(0, "analysis")
import metrics as M
import resampling as R
from funded_lab import with_slippage
from portfolio import load_mc

OUT = Path("analysis/outputs/codex_ldf_flow_sizing")
REPORT = Path("mc_reports/Backtesting Strategy Performance Report _ @NQ - 5 Minutes Late Day Flow.xlsx")
FUND_REF = "c23fc02"
END = pd.Timestamp("2022-12-31")
COEFFICIENTS = {"TQQQ": 6, "SQQQ": 12, "QLD": 2, "QID": 6}


def sha(data):
    return hashlib.sha256(data).hexdigest()


def write_json(path, value):
    path.write_text(json.dumps(value, indent=2, default=str, allow_nan=True) + "\n")


def git_blob(ref, path):
    return subprocess.check_output(["git", "show", f"{ref}:{path}"], stderr=subprocess.PIPE)


def truncate_grid():
    with open("scratch_data/g24.pkl", "rb") as f:
        original = pickle.load(f)
    g = copy.copy(original)
    n = int(np.searchsorted(original.ts, np.datetime64("2023-01-01")))
    for a in ("o", "h", "l", "c", "v", "ts", "dates", "last_j"):
        setattr(g, a, getattr(original, a)[:n])
    g.D = n
    g.trading = original.trading[original.trading < n]
    assert pd.Timestamp(g.ts[-1]) <= END
    return g


def fund_inputs_and_archive(sessions):
    a = pd.DataFrame(index=sessions)
    at = datetime.now(timezone.utc)
    folder = OUT / "fund_snapshots" / at.strftime("%Y%m%dT%H%M%SZ")
    folder.mkdir(parents=True, exist_ok=True)
    records = []
    for ticker in COEFFICIENTS:
        b = git_blob(FUND_REF, f"data/letf/{ticker}.csv")
        (folder / f"{ticker}.csv").write_bytes(b)
        x = pd.read_csv(io.BytesIO(b), usecols=["Date", "Assets Under Management"])
        x.index = pd.to_datetime(x.pop("Date"), format="%m/%d/%Y")
        assert not x.index.duplicated().any()
        a[ticker] = x["Assets Under Management"].reindex(sessions)
        records.append(dict(ticker=ticker, stored_at_utc=at.isoformat(), sha256=sha(b),
                            source_git=f"{FUND_REF}:data/letf/{ticker}.csv",
                            issuer_url=f"https://accounts.profunds.com/etfdata/ByFund/{ticker}-historical_nav.csv",
                            acquisition="Copy of Claude's supplied issuer snapshot; no independent fresh issuer retrieval",
                            independent_issuer_retrieval_time=None,
                            network_note="Fresh issuer probe blocked by environment HTTP proxy (403); no bypass attempted."))
    write_json(folder / "provenance.json", records)
    return a, str(folder)


def warmup_frame():
    p = Path("data/NQ_1min_2017.csv")
    if p.exists():
        b = p.read_bytes()
        source = str(p)
    else:
        source = "origin/claude/nq-trading-strategy-params-5h9d20:data/NQ_1min_2017.csv"
        try:
            b = git_blob("origin/claude/nq-trading-strategy-params-5h9d20", "data/NQ_1min_2017.csv")
        except subprocess.CalledProcessError:
            raise SystemExit("WAITING FOR PROMISED 2017 DAYTIME EXPORT; no inputs or outcomes evaluated")
    x = pd.read_csv(io.BytesIO(b), usecols=["date", "time", "close", "volume"])
    assert not x.duplicated(["date", "time"]).any(), "Duplicate warmup bars; stop before outcomes"
    d = pd.to_datetime(x.date.astype(str), format="%Y%m%d")
    assert (d.dt.year == 2017).all()
    assert x.time.between(800, 1515).all(), "Expected 08:00-15:15 CT daytime export"
    assert d.min() <= pd.Timestamp("2017-01-04") and d.max() >= pd.Timestamp("2017-12-28")
    return x, dict(source=source, sha256=sha(b), bars=len(x), first=str(d.min().date()), last=str(d.max().date()))


def entry_dates_only():
    x = pd.read_excel(REPORT, sheet_name="List of Trades", header=2, usecols=["Type", "Date"])
    x = x[x.Type.astype(str).str.startswith("Entry")].copy()
    x["date"] = pd.to_datetime(x.Date).dt.normalize()
    x = x[x.date.between("2018-01-01", END)].reset_index(drop=True)
    return x[["date"]]


def rank_feature(h):
    low = h.shift(1).rolling(252, min_periods=126).quantile(1/3, interpolation="linear")
    high = h.shift(1).rolling(252, min_periods=126).quantile(2/3, interpolation="linear")
    group = pd.Series("unavailable", index=h.index)
    ok = h.notna() & low.notna() & high.notna()
    group.loc[ok] = "middle"
    group.loc[ok & (h < low)] = "low"
    group.loc[ok & (h > high)] = "high"
    rank = pd.Series(np.nan, index=h.index)
    for i in np.flatnonzero(ok.to_numpy()):
        hist = h.iloc[max(0, i-252):i].dropna()
        rank.iloc[i] = ((hist < h.iloc[i]).sum() + .5*(hist == h.iloc[i]).sum()) / len(hist)
    return group, rank


def inputs():
    OUT.mkdir(parents=True, exist_ok=True)
    calendar = xc.get_calendar("XNYS", start="2017-01-01", end="2022-12-31")
    schedule = calendar.schedule
    sessions = pd.DatetimeIndex(schedule.index).tz_localize(None)
    assets, archive = fund_inputs_and_archive(sessions)
    warm, warm_manifest = warmup_frame()
    g = truncate_grid()
    day_lookup = pd.Series(np.arange(g.D), index=pd.DatetimeIndex(g.ts))
    wclose = warm.set_index(["date", "time"])["close"]
    wgroups = {date: block for date, block in warm.groupby("date")}
    frame = pd.DataFrame(index=sessions)
    frame.index.name = "date"
    frame["cash_close"] = np.nan
    frame["last_hour_notional"] = np.nan
    frame["last_hour_contracts"] = np.nan
    for i, day in enumerate(sessions):
        close_ts = schedule["close"].iloc[i].tz_convert("America/Chicago")
        close_j = close_ts.hour*60 + close_ts.minute
        regular = close_j == 900
        if day.year == 2017:
            date_int = int(day.strftime("%Y%m%d"))
            frame.loc[day, "cash_close"] = wclose.get((date_int, close_ts.hour*100+close_ts.minute), np.nan)
            z = wgroups.get(date_int)
            if regular and z is not None:
                z = z[z.time.between(1401, 1500)]
                px = z.close.to_numpy(float)
                vol = z.volume.to_numpy(float)
                if len(z) == 60 and np.isfinite(px).all() and np.isfinite(vol).all() and (px > 0).all() and vol.sum() > 0:
                    frame.loc[day, "last_hour_notional"] = float(np.sum(20*px*vol))
                    frame.loc[day, "last_hour_contracts"] = float(vol.sum())
        elif day in day_lookup.index:
            d = int(day_lookup.loc[day])
            frame.loc[day, "cash_close"] = g.c[d, close_j]
            px, vol = g.c[d, 841:901], g.v[d, 841:901]
            if regular and np.isfinite(px).all() and np.isfinite(vol).all() and (px > 0).all() and vol.sum() > 0:
                frame.loc[day, "last_hour_notional"] = float(np.sum(20*px*vol))
                frame.loc[day, "last_hour_contracts"] = float(vol.sum())
    frame["expected_notional"] = np.nan
    for i in range(len(frame)):
        history = frame.last_hour_notional.iloc[max(0, i-40):i].dropna().tail(20)
        if len(history) == 20:
            frame.iloc[i, frame.columns.get_loc("expected_notional")] = history.median()
    for lag, key in ((1, "R"), (2, "S")):
        k = sum(assets[t].shift(lag)*c for t, c in COEFFICIENTS.items())
        h = .01*k/frame.expected_notional
        frame[f"H_{key}"] = h
        frame[f"group_{key}"], frame[f"rank_{key}"] = rank_feature(h)
        frame[f"q_{key}"] = frame[f"group_{key}"].map({"low": 1, "middle": 2, "high": 3, "unavailable": 2})
    for lag in (63, 252):
        frame[f"return_{lag}"] = frame.cash_close.shift(1)/frame.cash_close.shift(lag+1)-1
    ready = (frame.group_R != "unavailable") & (frame.group_S != "unavailable") & frame[["return_63", "return_252"]].notna().all(axis=1)
    eligible_start = frame.index[ready & (frame.index >= "2018-01-01")][0]
    # Missing regression controls are explicit fixed-size fallback, not a silently selected sample.
    frame["joint_ready"] = ready
    frame.loc[~ready, ["q_R", "q_S"]] = 2
    entries = entry_dates_only().merge(frame.reset_index(), on="date", how="left", validate="many_to_one")
    entries["warmup_excluded"] = entries.date < eligible_start
    evaluated = entries[~entries.warmup_excluded]
    covered = evaluated[evaluated.joint_ready.fillna(False)]
    years = sorted(evaluated.date.dt.year.unique().tolist())
    counts = covered.groupby([covered.date.dt.year, "group_R"]).size().unstack(fill_value=0).reindex(years, fill_value=0)
    counts.to_csv(OUT / "input_entry_group_counts_by_year.csv")
    n_hi = int((covered.group_R == "high").sum())
    n_lo = int((covered.group_R == "low").sum())
    coverage = float(len(covered)/len(evaluated))
    both_every_year = bool(all(counts.get(k, pd.Series(0, index=years)).gt(0).all() for k in ("high", "low")))
    support = bool(coverage >= .90 and n_hi >= 100 and n_lo >= 100 and both_every_year)
    market = frame[(frame.index >= "2018-01-01") & ready]
    correlations = []
    for population, z in (("cash_sessions_2018_22", market), ("evaluated_LDF_entries", covered)):
        for lag in (63, 252):
            for method in ("pearson", "spearman"):
                correlations.append(dict(population=population, momentum_sessions=lag, method=method,
                                         correlation=float(z[f"rank_R"].corr(z[f"return_{lag}"], method=method)), observations=len(z)))
    pd.DataFrame(correlations).to_csv(OUT / "input_rank_momentum_correlations.csv", index=False)
    market.groupby("group_R")[["return_63", "return_252"]].mean().to_csv(OUT / "input_momentum_means_by_group.csv")
    market.groupby([market.index.year, "group_R"]).size().rename("cash_sessions").to_csv(OUT / "input_cash_group_counts_by_year.csv")
    is_roll = frame.index.month.isin([3, 6, 9, 12]) & (frame.index.day >= 5) & (frame.index.day <= 16) & frame.index.dayofweek.isin([3, 4])
    frame["roll_thursday_friday_window"] = is_roll
    roll = frame.loc[is_roll, ["last_hour_contracts", "last_hour_notional", "expected_notional"]].copy()
    roll["relative_to_lagged_expected_notional"] = roll.last_hour_notional/roll.expected_notional
    roll.to_csv(OUT / "input_roll_volume_diagnostics.csv")
    frame.to_csv(OUT / "input_cash_features_2017_22.csv")
    entries.to_csv(OUT / "input_entry_features.csv", index=False)
    manifest = dict(plan_amendment="6ca8950", source_funds=FUND_REF, warmup=warm_manifest,
                    evaluation_start=str(eligible_start.date()), evaluation_end="2022-12-31",
                    original_design_entries=len(entries), warmup_excluded_entries=int(entries.warmup_excluded.sum()),
                    evaluated_entries=len(evaluated), covered_entries=len(covered), joint_coverage=coverage,
                    high_entries=n_hi, low_entries=n_lo, support_gate_passed=support,
                    high_low_every_evaluated_year=both_every_year,
                    missing_market_close_sessions=[str(d.date()) for d in frame.index[frame.cash_close.isna()]],
                    implied_t2_gap_assuming_sd810=float(2*810*np.sqrt(1/n_hi+1/n_lo)) if min(n_hi, n_lo)>0 else None,
                    fund_snapshot_archive=archive, report_sha256=sha(REPORT.read_bytes()),
                    cache_sha256=sha(Path("scratch_data/g24.pkl").read_bytes()),
                    market_features_max_date=str(frame.index.max().date()),
                    candidate_strategy_returns_computed=False, return_trials=0,
                    raw_nq_input_hashes={str(p):sha(p.read_bytes()) for p in sorted(Path("data").glob("NQ_1min*20[12][0-9].csv")) if int(p.stem[-4:])<=2022})
    write_json(OUT / "input_checkpoint.json", manifest)
    print(json.dumps(manifest, indent=2))
    print(pd.DataFrame(correlations).to_string(index=False))
    print("RETURN-FREE INPUT CHECKPOINT COMPLETE. Design returns remain unopened.")


def hac_fit(y, x, contrast, lags=5):
    """OLS with Bartlett Newey-West covariance; no optional statsmodels dependency."""
    x, y = np.asarray(x, float), np.asarray(y, float)
    n, p = x.shape
    rank = int(np.linalg.matrix_rank(x))
    if rank != p or n <= p:
        return dict(estimate=np.nan, t=np.nan, rank=rank, columns=p, n=n, singular=True)
    bread = np.linalg.inv(x.T@x)
    beta = np.linalg.lstsq(x, y, rcond=None)[0]
    scores = x*(y-x@beta)[:, None]
    meat = scores.T@scores
    for lag in range(1, min(lags+1, n)):
        c = scores[lag:].T@scores[:-lag]
        meat += (1-lag/(lags+1))*(c+c.T)
    cov = bread@meat@bread*n/(n-p)
    contrast = np.asarray(contrast, float)
    estimate = float(contrast@beta)
    variance = float(contrast@cov@contrast)
    se = np.sqrt(max(0, variance))
    return dict(estimate=estimate, se=float(se), t=float(estimate/se) if se>0 else np.nan,
                rank=rank, columns=p, n=n, condition_number=float(np.linalg.cond(x)), singular=False)


def adjusted_contrast(t, key):
    z = t[t.joint_ready & t.regression_ready].copy()
    x = pd.DataFrame({"intercept": 1., "high": (z[f"group_{key}"]=="high").astype(float),
                      "low": (z[f"group_{key}"]=="low").astype(float)}, index=z.index)
    for col in ("absolute_rth_return", "move_atr", "atr_fraction", "return_63", "return_252", "side"):
        s = z[col].astype(float)
        x[col] = (s-s.mean())/s.std(ddof=0) if s.std(ddof=0)>0 else s-s.mean()
    for year in sorted(z.date.dt.year.unique())[1:]:
        x[f"year_{year}"] = (z.date.dt.year==year).astype(float)
    c = np.zeros(x.shape[1]); c[1], c[2] = 1, -1
    return hac_fit(z.net_nq, x, c)


def load_only_design_mc(g, start):
    """Date-filter before the accepted loader can calculate any strategy P&L."""
    wb = openpyxl.load_workbook(REPORT, read_only=True, data_only=True)
    ws = wb["List of Trades"]
    rows = ws.iter_rows(values_only=True)
    next(rows); next(rows); header = list(next(rows))
    type_i, date_i = header.index("Type"), header.index("Date")
    selected = []
    for row in rows:
        if not str(row[type_i]).startswith(("Entry", "Exit")) or not row[date_i]:
            continue
        date = pd.Timestamp(row[date_i]).normalize()
        if start <= date <= END:
            selected.append(row)
    settings = list(wb["Settings"].iter_rows(values_only=True))
    wb.close()
    # Temporary date-filtered input lets us reuse load_mc without changing shared code.
    with tempfile.TemporaryDirectory(prefix="mc_design_", dir=OUT) as folder:
        p = Path(folder)/"design_only.xlsx"
        new = openpyxl.Workbook(); trade = new.active; trade.title = "List of Trades"
        trade.append([None]); trade.append([None]); trade.append(header)
        for row in selected:
            trade.append(row)
        sw = new.create_sheet("Settings")
        for row in settings:
            sw.append(row)
        new.save(p)
        return load_mc(g, str(p))


def tercile(s):
    a, b = s.quantile([1/3, 2/3]).to_numpy()
    return pd.Series(np.where(s<a, 0, np.where(s>b, 2, 1)), index=s.index)


def design():
    checkpoint_path = OUT/"input_checkpoint.json"
    manifest = json.loads(checkpoint_path.read_text())
    assert not manifest["candidate_strategy_returns_computed"]
    assert manifest["report_sha256"] == sha(REPORT.read_bytes())
    if not manifest["support_gate_passed"]:
        write_json(OUT/"design_decision.json", dict(verdict="WATCH", reason="Registered input support gate failed; no outcome test",
                   return_trials=0, candidate_configurations_evaluated=[], replication_opened=False, funded_simulations=0))
        print("STOP: input support gate failed; no strategy return computed.")
        return
    assert not (OUT/"design_decision.json").exists(), "Design already evaluated; no silent reruns"
    features = pd.read_csv(OUT/"input_entry_features.csv", parse_dates=["date"])
    features = features[~features.warmup_excluded]
    start = pd.Timestamp(manifest["evaluation_start"])
    g = truncate_grid()
    t = load_only_design_mc(g, start)
    assert t.date.max() <= END and len(t) == manifest["evaluated_entries"]
    gross = (t.exit-t.entry)*t.side*20
    t = with_slippage(t)
    t["gross_nq"] = gross
    t["net_nq"] = (t.exit-t.entry)*t.side*20-7.25
    t["cost_nq"] = t.gross_nq-t.net_nq
    assert np.allclose(t.cost_nq, 17.25), "Unexpected LDF fill type/cost; do not silently change cost model"
    t = t.merge(features, on="date", how="left", validate="one_to_one")
    atr = g.daily_atr(15)
    d = t.day.to_numpy(int)
    rth_open, decision_close = g.o[d, 511], g.c[d, 840]
    t["absolute_rth_return"] = abs(decision_close/rth_open-1)
    t["move_atr"] = abs(decision_close-rth_open)/atr[d]
    t["atr_fraction"] = atr[d]/rth_open
    t["regression_ready"] = np.isfinite(t[["absolute_rth_return", "move_atr", "atr_fraction", "return_63", "return_252"]]).all(axis=1)
    assert t.loc[t.joint_ready, "regression_ready"].all(), "Missing frozen regression controls; stop without selecting another sample"
    day_index = pd.DatetimeIndex(g.ts[g.trading])
    day_index = day_index[(day_index>=start)&(day_index<=END)]
    daily = pd.DataFrame(index=day_index)
    daily.index.name = "date"
    for key, q in (("B", np.full(len(t), 2)), ("R", t.q_R.to_numpy()), ("S", t.q_S.to_numpy())):
        t[f"net_{key}"] = q*t.net_nq/10
        t[f"gross_{key}"] = q*t.gross_nq/10
        daily[key] = t.groupby("date")[f"net_{key}"].sum().reindex(day_index, fill_value=0.)
    daily["increment_R"] = daily.R-daily.B
    daily["increment_S"] = daily.S-daily.B
    t["increment_R"] = t.net_R-t.net_B
    t["increment_S"] = t.net_S-t.net_B
    rows = []
    for key in ("B", "R", "S"):
        for period, choose in [("design", np.ones(len(t), bool))]+[(str(y), t.date.dt.year.to_numpy()==y) for y in sorted(t.date.dt.year.unique())]:
            for side, side_pick in (("all", np.ones(len(t),bool)), ("long", t.side.to_numpy()>0), ("short", t.side.to_numpy()<0)):
                sub = t[choose & side_pick]
                days = day_index if period=="design" else day_index[day_index.year==int(period)]
                x = sub.groupby("date")[f"net_{key}"].sum().reindex(days, fill_value=0.)
                perf = M.perf(x, sub[f"net_{key}"])
                net = sub[f"net_{key}"]
                equivalent = 2*len(sub)/10 if key=="B" else sub[f"q_{key}"].sum()/10
                perf.update(configuration=key, period=period, side=side, total_net=float(net.sum()),
                            t_trade=float(net.mean()/net.std()*np.sqrt(len(net))) if len(net)>1 and net.std()>0 else np.nan,
                            gross_total=float(sub[f"gross_{key}"].sum()), costs_total=float((sub[f"gross_{key}"]-net).sum()),
                            mean_net_per_nq_equivalent=float(net.sum()/equivalent) if equivalent else np.nan,
                            mean_gross_per_nq_equivalent=float(sub[f"gross_{key}"].sum()/equivalent) if equivalent else np.nan)
                rows.append(perf)
    pd.DataFrame(rows).to_csv(OUT/"design_standard_metrics.csv", index=False)
    group_rows = []
    for key in ("R", "S"):
        for label, z in t.groupby(f"group_{key}"):
            x = z.net_nq
            group_rows.append(dict(configuration=key, group=label, entries=len(z), mean_net_nq=x.mean(), sd_net_nq=x.std(),
                                   total_net_nq=x.sum(), mean_gross_nq=z.gross_nq.mean()))
    pd.DataFrame(group_rows).to_csv(OUT/"design_flow_groups.csv", index=False)
    adjusted = {k: adjusted_contrast(t, k) for k in ("R", "S")}
    write_json(OUT/"design_adjusted_contrasts.json", adjusted)
    hac_increment = hac_fit(daily.increment_R, np.ones((len(daily),1)), [1.])
    boot = R.bootstrap(daily.increment_R, n=5000, block=20, seed=2026102201)
    boot.to_csv(OUT/"design_increment_bootstrap.csv")

    valid = t.joint_ready & t.regression_ready
    t["year"] = t.date.dt.year
    t["move_tercile"] = -1
    t["momentum63_tercile"] = -1
    for _, z in t[valid].groupby("year"):
        t.loc[z.index, "move_tercile"] = tercile(z.move_atr)
        t.loc[z.index, "momentum63_tercile"] = tercile(z.return_63)
    groups = [z.index.to_numpy() for _, z in t[valid].groupby(["year", "side", "move_tercile", "momentum63_tercile"])]
    cells = []
    for ix in groups:
        z = t.loc[ix]
        cells.append(dict(year=int(z.year.iloc[0]), side=int(z.side.iloc[0]), move_tercile=int(z.move_tercile.iloc[0]),
                          momentum63_tercile=int(z.momentum63_tercile.iloc[0]), entries=len(ix),
                          low=int((z.q_R==1).sum()), middle=int((z.q_R==2).sum()), high=int((z.q_R==3).sum()),
                          movable=bool(z.q_R.nunique()>1)))
    pd.DataFrame(cells).to_csv(OUT/"design_placebo_cells.csv", index=False)
    rng = np.random.default_rng(2026102202)
    q = t.q_R.to_numpy(float)
    date_position = pd.Series(np.arange(len(day_index)), index=day_index).reindex(t.date).to_numpy(int)
    null = []
    for run in range(2000):
        shuffled = q.copy()
        for ix in groups:
            shuffled[ix] = rng.permutation(q[ix])
        net = shuffled*t.net_nq.to_numpy()/10
        path = np.bincount(date_position, weights=net, minlength=len(day_index))
        null.append(dict(run=run, incremental_net=float(net.sum()-t.net_B.sum()), max_dd=R.max_dd(path)))
    placebo = pd.DataFrame(null)
    placebo.to_csv(OUT/"design_matched_placebo.csv", index=False)
    actual_delta = float(t.increment_R.sum())
    p = float((1+(placebo.incremental_net>=actual_delta).sum())/(1+len(placebo)))
    actual_dd = R.max_dd(daily.R)
    annual = daily.groupby(daily.index.year).sum()
    annual.to_csv(OUT/"design_annual_totals.csv")
    candidate_equivalents = t.q_R.sum()/10
    gross_per_equivalent = float(t.gross_R.sum()/candidate_equivalents)
    gates = {
        "G1_input_support": bool(manifest["support_gate_passed"]),
        "G2_net_cost_years": bool(t.net_R.sum()>0 and gross_per_equivalent>=1.5*17.25 and (annual.R>0).sum()>=3),
        "G3_incremental_net_significance": bool(actual_delta>0 and (annual.increment_R>0).sum()>=3 and hac_increment["t"]>=2 and boot.loc["annual","p5"]>0),
        "G4_momentum_adjusted_contrast": bool(adjusted["R"]["estimate"]>0 and adjusted["R"]["t"]>=2 and not adjusted["R"]["singular"]),
        "G5_matched_placebo": bool(p<=.05 and actual_dd<=placebo.max_dd.median()),
        "G6_two_session_stress": bool(t.increment_S.sum()>0 and adjusted["S"]["estimate"]>0 and not adjusted["S"]["singular"]),
    }
    decision = dict(verdict="WATCH" if all(gates.values()) else "REJECT", design_gate_passed=all(gates.values()),
                    gates=gates, primary_total_incremental_net=actual_delta, stress_total_incremental_net=float(t.increment_S.sum()),
                    primary_incremental_daily_hac=hac_increment, primary_adjusted_contrast=adjusted["R"], stress_adjusted_contrast=adjusted["S"],
                    placebo_p=p, primary_max_dd=actual_dd, placebo_median_max_dd=float(placebo.max_dd.median()),
                    gross_per_nq_equivalent=gross_per_equivalent, increment_annual_bootstrap_p5=float(boot.loc["annual","p5"]),
                    hypothesis_trials=1, nonbaseline_configurations_evaluated=["R", "S"], addition_evaluated=False,
                    replication_opened=False, funded_simulations=0, candidate_strategy_returns_max_date=str(t.date.max().date()),
                    input_checkpoint_sha256=sha(checkpoint_path.read_bytes()), exploratory_snapshot_assumption=True)
    t.to_csv(OUT/"design_trades.csv", index=False)
    daily.to_csv(OUT/"design_daily_pnl.csv")
    write_json(OUT/"design_decision.json", decision)
    print(json.dumps(decision, indent=2))
    print("STOP AFTER DESIGN. No 2023+ candidate calculations or funded-account simulation.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("phase", choices=["inputs", "design"])
    args = parser.parse_args()
    inputs() if args.phase=="inputs" else design()
